# -*- coding: utf-8 -*-
"""Génère le PowerPoint de présentation Canal Cup 2026 (16:9, charte Canal+)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ─── Charte ────────────────────────────────────────────────────────────────
BLACK   = RGBColor(0x0A, 0x0A, 0x0A)
PANEL   = RGBColor(0x16, 0x16, 0x16)
CARD    = RGBColor(0x1E, 0x1E, 0x1E)
YELLOW  = RGBColor(0xFF, 0xDD, 0x00)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GRAY    = RGBColor(0x9A, 0x9A, 0x9A)
LGRAY   = RGBColor(0xC8, 0xC8, 0xC8)
RED     = RGBColor(0xE5, 0x3E, 0x3E)
GREEN   = RGBColor(0x33, 0xCC, 0x77)

FONT = "Segoe UI"
FONT_BLACK = "Segoe UI Black"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BLACK
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def setp(p, text, size, color, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         space_after=6, italic=False):
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return p


def rect(s, x, y, w, h, fill, line=None, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def kicker_title(s, kicker, title="", accent=YELLOW):
    h = 0.95 if title else 0.45
    rect(s, 0.7, 0.62, 0.16, h, accent, shape=MSO_SHAPE.RECTANGLE)
    _, tf = box(s, 1.0, 0.55, 11.6, 1.3)
    setp(tf.paragraphs[0], kicker.upper(), 13, accent, bold=True, space_after=2)
    if title:
        setp(tf.add_paragraph(), title, 30, WHITE, bold=True, font=FONT_BLACK)


def footer(s, idx, total):
    _, tf = box(s, 0.7, 6.95, 12.0, 0.4)
    p = tf.paragraphs[0]
    setp(p, "CANAL CUP 2026", 9, GRAY, bold=True, space_after=0)
    _, tf2 = box(s, 11.6, 6.95, 1.1, 0.4)
    setp(tf2.paragraphs[0], f"{idx:02d} / {total:02d}", 9, GRAY, align=PP_ALIGN.RIGHT, space_after=0)


def bullets(s, items, x=1.0, y=2.0, w=11.4, gap=8, size=16):
    _, tf = box(s, x, y, w, 4.6)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(it, tuple):
            head, sub = it
            p.space_after = Pt(2)
            r = p.add_run(); r.text = "▸  "; r.font.size = Pt(size); r.font.bold = True
            r.font.color.rgb = YELLOW; r.font.name = FONT
            r2 = p.add_run(); r2.text = head; r2.font.size = Pt(size); r2.font.bold = True
            r2.font.color.rgb = WHITE; r2.font.name = FONT
            ps = tf.add_paragraph(); ps.space_after = Pt(gap)
            rs = ps.add_run(); rs.text = "      " + sub; rs.font.size = Pt(13)
            rs.font.color.rgb = LGRAY; rs.font.name = FONT
        else:
            p.space_after = Pt(gap)
            r = p.add_run(); r.text = "▸  "; r.font.size = Pt(size); r.font.bold = True
            r.font.color.rgb = YELLOW; r.font.name = FONT
            r2 = p.add_run(); r2.text = it; r2.font.size = Pt(size)
            r2.font.color.rgb = LGRAY; r2.font.name = FONT


TOTAL = 18
n = 0

def F(s):
    global n
    footer(s, n, TOTAL)

# ─── 1. Titre ────────────────────────────────────────────────────────────────
n += 1
s = slide()
rect(s, 0, 0, 13.333, 0.18, YELLOW, shape=MSO_SHAPE.RECTANGLE)
_, tf = box(s, 1.0, 2.3, 11.3, 2.2)
p = tf.paragraphs[0]; p.space_after = Pt(0)
r = p.add_run(); r.text = "CANAL "; r.font.size = Pt(66); r.font.bold = True; r.font.name = FONT_BLACK; r.font.color.rgb = YELLOW
r = p.add_run(); r.text = "CUP "; r.font.size = Pt(66); r.font.bold = True; r.font.name = FONT_BLACK; r.font.color.rgb = WHITE
r = p.add_run(); r.text = "2026"; r.font.size = Pt(66); r.font.bold = True; r.font.name = FONT_BLACK; r.font.color.rgb = GRAY
setp(tf.add_paragraph(), "Vivre la Coupe du Monde ensemble, au bureau.", 22, LGRAY, space_after=4)
_, tf2 = box(s, 1.0, 4.5, 11.0, 1.0)
setp(tf2.paragraphs[0], "Pronostics • Équipes • Babyfoot • Quiz • Animations RSE • Bonne ambiance", 15, YELLOW, bold=True)
setp(tf2.add_paragraph(), "Une expérience interne pensée pour tout le monde — même celles et ceux qui détestent le foot.", 13, GRAY, italic=True)
F(s)

# ─── 2. Le constat ───────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Le constat", "Un grand événement, et chacun dans son coin")
bullets(s, [
    ("La Coupe du Monde 2026, c'est 1 mois de matchs.", "Un moment fédérateur rare… qui d'habitude se vit en silence devant un écran."),
    ("Tout le monde n'aime pas le foot.", "Les conversations « foot » excluent vite une partie des équipes."),
    ("La RSE interne manque souvent de ce qui rassemble vraiment.", "On veut du lien, du fun, de la transversalité entre services."),
    ("Et si l'événement devenait un terrain de jeu commun ?", "Accessible, taquin, par équipes — où l'ambiance compte autant que le score."),
], y=1.9)
F(s)

# ─── 3. La vision ─────────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "La vision", "Canal Cup : la Coupe du Monde, version Canal+")
_, tf = box(s, 1.0, 1.85, 11.4, 1.2)
setp(tf.paragraphs[0], "Une web-app mobile qui transforme le mois de Coupe du Monde en compétition interne par équipes.", 18, WHITE, bold=True, space_after=6)
setp(tf.add_paragraph(), "On pronostique, on joue au quiz, au babyfoot, on relève des défis RSE — et chaque point fait gagner son équipe.", 15, LGRAY)
cards = [
    ("EN BINÔMES", "Autant d'équipes\nque de duos."),
    ("5 FAÇONS\nDE MARQUER", "Foot, quiz, babyfoot,\ndéfis, social."),
    ("100% INCLUSIF", "Zéro foot requis\npour gagner."),
    ("UN RITUEL", "Matinale IA + Mode TV\nqui anime la salle."),
]
cx = 1.0
for head, sub in cards:
    c = rect(s, cx, 4.0, 2.78, 2.4, CARD, line=YELLOW, line_w=1.0)
    tfc = c.text_frame; tfc.word_wrap = True; tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
    tfc.margin_left = Inches(0.18); tfc.margin_right = Inches(0.18)
    setp(tfc.paragraphs[0], head, 16, YELLOW, bold=True, align=PP_ALIGN.CENTER, font=FONT_BLACK, space_after=6)
    setp(tfc.add_paragraph(), sub, 12, LGRAY, align=PP_ALIGN.CENTER)
    cx += 2.97
F(s)

# ─── 4. Les équipes ───────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Le coeur du jeu", "Tout se joue en équipe")
bullets(s, [
    ("On forme un binôme Canal Cup.", "2 personnes max, un capitaine + un code d'invitation. Autant d'équipes que de duos — c'est l'unité de compétition."),
    ("Chaque point gagné alimente son équipe.", "Pronostics, quiz, babyfoot, défis : tout se consolide au niveau du binôme."),
    ("Mais les stats perso restent à soi.", "Tes séries, tes scores exacts : ta fierté individuelle, sans pression de groupe."),
    ("On ne vote pas pour soi… mais on chambre les autres avec style.", "La rivalité bon enfant fait partie du jeu."),
], y=1.9)
F(s)

# ─── 5. Les 5 piliers de score ───────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Le moteur", "5 façons de marquer des points")
pillars = [
    ("PRONOSTICS", "35%", "Score exact +10, bon résultat +5. Pondéré par phase : finale ×3.", YELLOW),
    ("ANIMATIONS RSE", "25%", "Défis physiques en salle. Participer = minimum 5 pts garantis.", GREEN),
    ("QUIZ LIVE", "20%", "Bonne réponse en < 5s : +5 pts. L'énergie du direct.", WHITE),
    ("BABYFOOT", "20%", "Tournoi de table. Chaque victoire rapporte à l'équipe.", WHITE),
    ("SOCIAL (REVIVEZ)", "+", "Likes & moments cultes : métrique sociale, hors classement.", GRAY),
]
y = 1.95
for name, pct, desc, col in pillars:
    rect(s, 1.0, y, 1.5, 0.78, PANEL, shape=MSO_SHAPE.RECTANGLE)
    _, tfp = box(s, 1.0, y+0.06, 1.5, 0.7)
    setp(tfp.paragraphs[0], pct, 26, col, bold=True, align=PP_ALIGN.CENTER, font=FONT_BLACK, space_after=0)
    _, tfn = box(s, 2.75, y-0.02, 9.6, 0.85)
    setp(tfn.paragraphs[0], name, 16, col, bold=True, space_after=1)
    setp(tfn.add_paragraph(), desc, 12, LGRAY)
    y += 0.92
F(s)

# ─── 6. Pronostics ────────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · Pronostics")
_, tf = box(s, 1.0, 1.45, 11.4, 0.6)
setp(tf.paragraphs[0], "Pronostique chaque match — même sans rien connaître au foot", 22, WHITE, bold=True)
bullets(s, [
    ("Tu donnes ton score, ça se sauvegarde tout seul.", "Score exact +10 · bon résultat +5 · bonne différence +3."),
    ("Les « tendances bureau » te montrent ce que parient les autres.", "% sur victoire / nul / défaite : tu suis la foule si tu hésites."),
    ("Des bonus longue durée pour les ambitieux.", "Vainqueur de la Coupe (+20), meilleur buteur (+10), série parfaite (+5)."),
    ("Plus on avance, plus ça compte.", "Phase de groupes ×1 → finale ×3. Le suspense monte jusqu'au bout."),
], y=2.15)
F(s)

# ─── 7. Quiz Live ─────────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · Quiz Live")
_, tf = box(s, 1.0, 1.45, 11.4, 0.6)
setp(tf.paragraphs[0], "Le quiz en direct qui met le feu à la salle", 22, WHITE, bold=True)
bullets(s, [
    ("L'animateur lance, tout le monde répond en même temps.", "Même question, même chrono synchronisé serveur — pas de triche possible."),
    ("20 secondes, 4 réponses, et le bonus rapidité.", "Bonne réponse en moins de 5s : +5 pts. Plus lent : +3 pts."),
    ("Des questions variées : foot, culture, Canal+, général.", "Tout le monde a sa chance, pas seulement les experts."),
    ("Un mode « écran géant » pour la projection.", "Game-show, gros boutons colorés, révélation spectaculaire."),
], y=2.15)
F(s)

# ─── 8. Babyfoot + Animations ────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Features · Babyfoot & Animations RSE", "Le jeu sort de l'écran")
# deux colonnes
rect(s, 1.0, 1.95, 5.55, 4.4, CARD, line=YELLOW, line_w=1.0)
_, tf1 = box(s, 1.3, 2.15, 5.0, 4.1)
setp(tf1.paragraphs[0], "BABYFOOT", 18, YELLOW, bold=True, font=FONT_BLACK, space_after=4)
setp(tf1.add_paragraph(), "« Moins de VAR, plus de chaos. »", 13, GRAY, italic=True, space_after=10)
for t in ["Tournoi de table entre les binômes.",
          "Chaque victoire rapporte des points au classement général.",
          "Le foot le plus inclusif qui soit : tout le monde sait jouer au baby."]:
    setp(tf1.add_paragraph(), "▸ " + t, 13.5, LGRAY, space_after=8)
rect(s, 6.78, 1.95, 5.55, 4.4, CARD, line=GREEN, line_w=1.0)
_, tf2 = box(s, 7.08, 2.15, 5.0, 4.1)
setp(tf2.paragraphs[0], "ANIMATIONS RSE", 18, GREEN, bold=True, font=FONT_BLACK, space_after=4)
setp(tf2.add_paragraph(), "Les défis qui rythment la Coupe.", 13, GRAY, italic=True, space_after=10)
for t in ["Défis physiques en salle : blind test, penalty, photo, quiz d'équipe…",
          "Tu te présentes à l'animateur, il valide, ça crédite ton équipe.",
          "Participer = au moins 5 points garantis. Ici, l'effort compte.",
          "Solo ou en groupe (points partagés)."]:
    setp(tf2.add_paragraph(), "▸ " + t, 13.5, LGRAY, space_after=7)
F(s)

# ─── 9. Classement & médailles ────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · Classement")
_, tf = box(s, 1.0, 1.45, 11.4, 0.6)
setp(tf.paragraphs[0], "Un classement live… et des médailles absurdes", 22, WHITE, bold=True)
bullets(s, [
    ("Toutes les équipes au classement, mis à jour après chaque match.", "🥇🥈🥉 pour le podium, puis chaque binôme classé + le détail des points par pilier."),
    ("Des « médailles absurdes » qui dédramatisent.", "Pire prédicteur de la semaine, match le plus mal pronostiqué… on rit de soi."),
    ("La compétition reste légère et taquine.", "Le but, c'est l'ambiance — pas d'écraser le voisin."),
], y=2.15)
F(s)

# ─── 10. Revivez ──────────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · Revivez")
_, tf = box(s, 1.0, 1.45, 11.4, 0.6)
setp(tf.paragraphs[0], "Les Archives du VAR : fails, phrases cultes, moments de gloire", 21, WHITE, bold=True)
bullets(s, [
    ("Un mur social des meilleurs moments de la Coupe Cup.", "Fails 💥 · phrases cultes 💬 · photos 📸 · exploits babyfoot 🏓."),
    ("On vote pour les meilleurs — sans algorithme toxique.", "Les plus drôles remontent. On lit, on vote, on rit."),
    ("Chaque vote reçu = +1 point social pour l'équipe autrice.", "Une couche d'engagement qui n'a rien à voir avec le foot."),
], y=2.15)
F(s)

# ─── 11. Matinale IA ──────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · La Matinale", "Le rituel quotidien, écrit par l'IA", accent=YELLOW)
bullets(s, [
    ("Chaque matin, un brief unique généré par l'IA (Gemini).", "Résumé de la veille, enjeux du jour, classement — comme un édito sportif."),
    ("Le « fail du jour » : le moment honteux à ne pas rater.", "De quoi lancer les conversations à la machine à café."),
    ("Une voix qui anime la compétition, jour après jour.", "Le ton est taquin, bon enfant — jamais robotique."),
    ("Visible dans l'app, sur l'accueil et en Mode TV.", "Un rendez-vous fixe qui crée l'habitude."),
], y=1.95)
F(s)

# ─── 12. Mode TV ──────────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Feature · Mode TV", "L'écran de la salle devient un spectacle")
bullets(s, [
    ("Un diaporama plein écran pour beamer, TV de couloir ou écran de salle.", "Les slides tournent automatiquement — zéro manipulation."),
    ("Scores en direct, classement, duel du tournoi, animations en cours…", "Données rafraîchies en continu. La salle vit au rythme de la Coupe."),
    ("Un QR code en permanence pour rejoindre l'app.", "On scanne, on est dans le jeu en quelques secondes."),
    ("Et l'IA commente l'ambiance en direct.", "Bandeau qui pulse selon la tension des matchs."),
], y=1.95)
F(s)

# ─── 13. APERÇU SCORE EN DIRECT (maquette) ─────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Aperçu", "Le bandeau « Score en direct »")
# carte score
cardx, cardy, cardw, cardh = 2.35, 1.95, 8.65, 3.55
rect(s, cardx, cardy, cardw, cardh, PANEL, line=YELLOW, line_w=1.5)
# badge live + minute
liveb = rect(s, cardx+0.35, cardy+0.32, 2.0, 0.5, RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
ltf = liveb.text_frame; ltf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(ltf.paragraphs[0], "● EN DIRECT  ·  67'", 13, WHITE, bold=True, align=PP_ALIGN.CENTER, space_after=0)
_, chtf = box(s, cardx+cardw-3.2, cardy+0.34, 2.85, 0.45)
setp(chtf.paragraphs[0], "Canal+  •  beIN Sports", 12, GRAY, align=PP_ALIGN.RIGHT, bold=True, space_after=0)
# équipe A
_, ta = box(s, cardx+0.3, cardy+1.05, 3.0, 1.6)
setp(ta.paragraphs[0], "🇫🇷", 40, WHITE, align=PP_ALIGN.CENTER, space_after=0)
setp(ta.add_paragraph(), "FRANCE", 20, WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_BLACK)
# score
_, sc = box(s, cardx+3.1, cardy+1.15, 2.45, 1.5)
setp(sc.paragraphs[0], "2  -  1", 54, YELLOW, bold=True, align=PP_ALIGN.CENTER, font=FONT_BLACK, space_after=0)
# équipe B
_, tb = box(s, cardx+cardw-3.3, cardy+1.05, 3.0, 1.6)
setp(tb.paragraphs[0], "🇧🇷", 40, WHITE, align=PP_ALIGN.CENTER, space_after=0)
setp(tb.add_paragraph(), "BRÉSIL", 20, WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_BLACK)
# buteurs
_, bz = box(s, cardx+0.3, cardy+2.7, cardw-0.6, 0.55)
setp(bz.paragraphs[0], "⚽ Mbappé 23', 54'                    ⚽ Vinícius Jr 61'", 13, LGRAY, align=PP_ALIGN.CENTER, space_after=0)
# bandeau tendances pronos sous la carte
trec = rect(s, cardx, cardy+cardh+0.18, cardw, 0.62, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
ttf = trec.text_frame; ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
setp(ttf.paragraphs[0], "Tendance pronos du bureau :   58% France   ·   22% nul   ·   20% Brésil",
     13, YELLOW, bold=True, align=PP_ALIGN.CENTER, space_after=0)
# légende
_, lg = box(s, 1.0, 6.35, 11.4, 0.55)
setp(lg.paragraphs[0], "Maquette indicative — le bandeau Score en direct, tel qu'il apparaît en Mode TV et sur l'accueil de l'app.",
     12, GRAY, italic=True, align=PP_ALIGN.CENTER)
F(s)

# ─── 14. Accessible aux non-footeux ────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Le parti pris", "Conçu pour celles et ceux qui n'aiment pas le foot")
_, tf = box(s, 1.0, 1.85, 11.4, 0.7)
setp(tf.paragraphs[0], "Dès l'inscription, tu choisis ton niveau : Expert ⚽ · Amateur 📺 · ou « Ambiance » 🎉.", 17, WHITE, bold=True)
setp(tf.add_paragraph(), "« Le foot ? Je viens pour les petits fours et l'équipe. » — et c'est totalement OK.", 14, GRAY, italic=True)
bullets(s, [
    ("On peut gagner sans regarder un seul match.", "Quiz, babyfoot, défis RSE, moments cultes : 65% des points sont hors-pronostics."),
    ("Aucun jargon, aucun gatekeeping.", "Les tendances et les bonus rendent même les pronos accessibles à tous."),
    ("L'objectif est le lien, pas l'expertise.", "Transversalité entre services, fun, et fierté d'équipe."),
], y=3.0)
F(s)

# ─── 15. Comment c'est construit (IA) ──────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Et aussi…", "Construit autrement : l'IA comme copilote")
_, tf = box(s, 1.0, 1.9, 11.4, 1.0)
setp(tf.paragraphs[0], "Canal Cup a été conçu en m'appuyant fortement sur des IA génératives (Claude, Gemini) comme partenaires de développement.",
     17, WHITE, bold=True, space_after=6)
setp(tf.add_paragraph(), "La vision produit, l'architecture, les arbitrages UX et la cohérence d'ensemble : ça, c'est humain.", 14, LGRAY)
bullets(s, [
    ("Moi : le concept, l'expérience, le système de points RSE, les choix produit.", "Du mode TV au scoring par équipes — chaque décision est dirigée."),
    ("L'IA : un accélérateur de conception, de prototypage et de code.", "Une équipe technique augmentée, pas un bouton magique."),
    ("Le résultat : un vrai MVP en production.", "Hébergé sur Vercel, base Supabase, vraie auth, données de matchs en direct."),
    ("Ce que ça dit pour Canal+ : on peut créer vite des expériences internes innovantes.", "C'est sans doute l'enseignement le plus stratégique du projet."),
], y=3.05, gap=6, size=15)
F(s)

# ─── 16. Statut & next ─────────────────────────────────────────────────────────
n += 1
s = slide()
kicker_title(s, "Où on en est", "Un MVP qui tourne, prêt à être joué")
bullets(s, [
    ("L'app est fonctionnelle et en ligne.", "Pronostics, quiz, classement, mode TV, matinale IA : tout est branché."),
    ("Synchro des vrais matchs de la Coupe du Monde 2026.", "Calendrier, scores, compositions — données réelles."),
    ("Prochaines étapes : tester en conditions réelles, animer, ajuster.", "Recueillir vos retours pour la version « jour J »."),
    ("Il ne manque qu'une chose : vous.", "Plus on est d'équipes, plus c'est drôle."),
], y=1.95)
F(s)

# ─── 17. Call to action ────────────────────────────────────────────────────────
n += 1
s = slide()
rect(s, 0, 0, 13.333, 0.18, YELLOW, shape=MSO_SHAPE.RECTANGLE)
_, tf = box(s, 1.0, 2.5, 11.3, 2.5)
setp(tf.paragraphs[0], "Prêts à jouer ?", 52, WHITE, bold=True, font=FONT_BLACK, space_after=10)
setp(tf.add_paragraph(), "Scannez le QR, formez votre binôme, et que le meilleur duo gagne.", 20, YELLOW, bold=True, space_after=6)
setp(tf.add_paragraph(), "Des binômes. Un seul gagnant. Beaucoup de drama.", 16, GRAY, italic=True)
F(s)

# ─── 18. Merci ─────────────────────────────────────────────────────────────────
n += 1
s = slide()
_, tf = box(s, 1.0, 3.0, 11.3, 1.6)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(8)
r = p.add_run(); r.text = "MERCI"; r.font.size = Pt(60); r.font.bold = True; r.font.name = FONT_BLACK; r.font.color.rgb = YELLOW
setp(tf.add_paragraph(), "Canal Cup 2026 — vivons la Coupe du Monde ensemble.", 18, LGRAY, align=PP_ALIGN.CENTER)
F(s)

out = r"C:\Users\vtrouillat\Documents\CanalCup\presentation\Canal_Cup_2026.pptx"
prs.save(out)
print("OK ->", out, "| slides:", len(prs.slides._sldIdLst))
